import openpyxl

#read data from an excel
'''workbook = openpyxl.load_workbook('utils/Book 1.xlsx')
sheet = workbook['Sheet1']
rows = sheet.max_row
cols = sheet.max_column

for rows in range(1,rows+1):
    for cols in range(1,cols+1):
        print(sheet.cell(row=rows,column=cols).value , end =" ")
        print()


#write data into an excel
workbook = openpyxl.load_workbook('utils/Book 1.xlsx')
sheet = workbook['Sheet1']
sheet.cell(row = 3 ,column = 7).value ="Playwright"

workbook.save('utils/Book 1.xlsx')'''


class ExcelUtility:
    def getRowCount(self,filename,sheetname):
        workbook = openpyxl.load_workbook(filename)
        sheet = workbook[sheetname]
        return sheet.max_row

    def getColumnCount(self,filename,sheetname):
        workbook = openpyxl.load_workbook(filename)
        sheet = workbook[sheetname]
        return sheet.max_column

    def getCellData(self,filename,sheetname,rowNumber,columnNumber):
        workbook = openpyxl.load_workbook(filename)
        sheet = workbook[sheetname]
        return sheet.cell(row = rowNumber,column=columnNumber).value

    def setCellData(self,filename, sheetname, rowNumber, columnNumber,inputData):
        workbook = openpyxl.load_workbook(filename)
        sheet = workbook[sheetname]
        sheet.cell(row=rowNumber, column=columnNumber).value = inputData
        workbook.save(filename)

    #print(getRowCount("Book1.xlsx","Sheet1"))



